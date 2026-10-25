"""Export the final paper to DOCX; external rendering is verified separately."""
import argparse
from pathlib import Path
import re
import xml.etree.ElementTree as ET
import markdown
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def text_run(paragraph, text, bold=False, italic=False, code=False):
    if not text:
        return
    run=paragraph.add_run(text)
    run.bold=bold;run.italic=italic
    if code:
        run.font.name='Consolas';run.font.size=Pt(9)


def inline(paragraph, element, base, bold=False, italic=False, code=False):
    text_run(paragraph,element.text,bold,italic,code)
    for child in element:
        if child.tag=='br':paragraph.add_run().add_break()
        elif child.tag=='img':
            path=(base/child.attrib['src']).resolve()
            with Image.open(path) as img:w,h=img.size
            width=min(16.2,18.0*w/h)
            paragraph.add_run().add_picture(str(path),width=Cm(width))
            paragraph.alignment=1
        elif child.tag=='a' and child.attrib.get('href','').startswith(('https://','http://')):
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),paragraph.part.relate_to(child.attrib['href'],RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');props=OxmlElement('w:rPr');style=OxmlElement('w:rStyle');style.set(qn('w:val'),'Hyperlink');props.append(style);run.append(props)
            value=OxmlElement('w:t');value.text=''.join(child.itertext());run.append(value);link.append(run);paragraph._p.append(link)
        elif child.tag=='a':
            text_run(paragraph,''.join(child.itertext()),bold,italic,code)
        else:inline(paragraph,child,base,bold or child.tag=='strong',italic or child.tag=='em',code or child.tag=='code')
        text_run(paragraph,child.tail,bold,italic,code)


def render(doc, element, base):
    tag=element.tag
    if re.fullmatch('h[1-6]',tag):
        p=doc.add_heading(level=max(1,min(3,int(tag[1])-1)));inline(p,element,base)
    elif tag=='p':
        p=doc.add_paragraph();inline(p,element,base)
        image=element.find('img')
        if image is not None:
            p.paragraph_format.keep_with_next=True
            caption=doc.add_paragraph(image.attrib.get('alt',''),style='Caption');caption.alignment=1
    elif tag=='pre':
        p=doc.add_paragraph(style='Code Block');p.add_run(''.join(element.itertext()).rstrip())
    elif tag=='table':
        rows=element.findall('.//tr');columns=max(len(list(row)) for row in rows)
        table=doc.add_table(rows=0,cols=columns);table.style='Table Grid'
        for index,row in enumerate(rows):
            cells=table.add_row().cells
            for cell,content in zip(cells,row):
                inline(cell.paragraphs[0],content,base,bold=index==0)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(4)
                    if len(rows)<=15:p.paragraph_format.keep_with_next=index<len(rows)-1
                    for run in p.runs:run.font.size=Pt(9 if columns>=5 else 10)
            properties=table.rows[-1]._tr.get_or_add_trPr()
            properties.append(OxmlElement('w:cantSplit'))
            if index==0:
                properties.append(OxmlElement('w:tblHeader'))
                for cell in cells:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF1E8');cell._tc.get_or_add_tcPr().append(shade)
        doc.add_paragraph()
    elif tag in ('ul','ol'):
        for item in element.findall('li'):
            p=doc.add_paragraph(style='List Number' if tag=='ol' else 'List Bullet');inline(p,item,base)
    elif tag=='blockquote':
        for child in element:
            p=doc.add_paragraph(style='Quote');inline(p,child,base)
    else:
        for child in element:render(doc,child,base)


def export(source,output):
    source=Path(source).resolve();output=Path(output).resolve();doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.3);sec.right_margin=Cm(2)
    sec.different_first_page_header_footer=True
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(6)
    for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        doc.styles[name].font.name='Times New Roman'
        fonts=doc.styles[name].element.get_or_add_rPr().rFonts
        if fonts is not None:
            for key in list(fonts.attrib):
                if 'Theme' in key:del fonts.attrib[key]
    doc.styles['Title'].font.size=Pt(26);doc.styles['Title'].font.bold=True
    doc.styles['Title'].font.color.rgb=RGBColor.from_string('214A3A')
    doc.styles['Subtitle'].font.size=Pt(14);doc.styles['Subtitle'].font.italic=False
    doc.styles['Subtitle'].font.color.rgb=RGBColor.from_string('263E32')
    for name in ['TOC 1','TOC 2']:
        style=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,1)
        style.base_style=normal;style.font.name='Times New Roman';style.font.size=Pt(10.5)
        style.paragraph_format.space_after=Pt(1);style.paragraph_format.space_before=Pt(0)
        style.paragraph_format.line_spacing=1;style.paragraph_format.keep_with_next=False
    for name in ['Heading 1','Heading 2','Heading 3']:
        doc.styles[name].font.color.rgb=RGBColor.from_string('214A3A');doc.styles[name].paragraph_format.keep_with_next=True
    code=doc.styles.add_style('Code Block',1);code.font.name='Consolas';code.font.size=Pt(8)
    code.paragraph_format.line_spacing=1;code.paragraph_format.keep_together=True
    header=sec.header.paragraphs[0];header.add_run('CAFELINK | BÀI TẬP LỚN LẬP TRÌNH MẠNG').font.size=Pt(9)
    footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('CafeLink V9.3 | ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.add_paragraph('BÀI TẬP LỚN MÔN LẬP TRÌNH MẠNG',style='Subtitle')
    p=doc.add_paragraph('CAFELINK',style='Title');p.paragraph_format.space_before=Pt(60)
    doc.add_paragraph('Hệ thống gọi món và điều phối phục vụ tại quán café theo mô hình Client/Server',style='Subtitle')
    doc.add_paragraph('Thiết kế TCP application protocol và phục hồi theo ngữ nghĩa/phạm vi phiên')
    doc.add_paragraph('Mốc phần mềm V9.3 · Schema 11\nChuẩn bị bàn giao cho hạn 31/10/2026')
    doc.add_paragraph('Thông tin nhóm, lớp, giảng viên và phân công: chờ nhóm bổ sung trước khi nộp.')
    doc.add_paragraph('Nghiệm thu thiết bị di động và kiểm tra Compilatio chưa thực hiện.')
    for paragraph in doc.paragraphs:paragraph.alignment=1
    doc.add_page_break();doc.add_paragraph('MỤC LỤC',style='Title')
    p=doc.add_paragraph()
    for kind in ['begin','separate','end']:
        if kind=='separate':
            run=OxmlElement('w:r');instruction=OxmlElement('w:instrText');instruction.set(qn('xml:space'),'preserve');instruction.text=' TOC \\o "1-2" \\h \\z \\u ';run.append(instruction);p._p.append(run)
        run=OxmlElement('w:r');fld=OxmlElement('w:fldChar');fld.set(qn('w:fldCharType'),kind);run.append(fld);p._p.append(run)
    doc.add_page_break()
    html=markdown.markdown(source.read_text(encoding='utf-8-sig'),extensions=['tables','fenced_code'],output_format='xhtml')
    root=ET.fromstring('<root>'+html+'</root>')
    for item in root:
        if item.tag=='h1':continue
        render(doc,item,source.parent)
    doc.core_properties.title='CafeLink — Báo cáo cuối môn Lập trình mạng'
    doc.core_properties.author='CafeLink — thông tin nhóm chờ bổ sung'
    doc.core_properties.subject='V9.3, TCP Client/Server, semantic recovery, experiments'
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
    reopened=Document(output)
    print({'docx':str(output),'paragraphs':len(reopened.paragraphs),'tables':len(reopened.tables),'images':len(reopened.inline_shapes)})


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--source',required=True);parser.add_argument('--output',required=True)
    args=parser.parse_args();export(args.source,args.output)
