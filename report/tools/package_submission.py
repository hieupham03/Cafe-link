"""Allowlisted source/delivery export. Never copy live demo data or developer caches."""
import argparse,hashlib,json,shutil,zipfile,re,unicodedata
from pathlib import Path
from datetime import datetime,timezone

ROOT=Path(__file__).resolve().parents[2]
IGNORE=shutil.ignore_patterns('node_modules','target','.tools','dist','experiment-dist','test-results','__pycache__','.git','history','*.pyc','*.log')

def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest().upper()
def copy_source(destination):
    destination=Path(destination).resolve()
    if destination.exists():raise FileExistsError('Use a fresh source destination')
    if not destination.is_relative_to(ROOT/'dist'):raise ValueError('Source export must stay in workspace dist')
    project=destination/'cafelink-system';project.mkdir(parents=True)
    for name in ['src','scripts','tools','release','web','docs']:
        shutil.copytree(ROOT/'cafelink-system'/name,project/name,ignore=IGNORE)
    for name in ['README.md','pom.xml','.gitignore']:shutil.copy2(ROOT/'cafelink-system'/name,project/name)
    (destination/'dist').mkdir();shutil.copy2(ROOT/'dist/CafeLink-Original-1.0.0-Baseline.zip',destination/'dist')
    shutil.copytree(ROOT/'.github',destination/'.github')
    shutil.copytree(ROOT/'final-project',destination/'final-project',ignore=shutil.ignore_patterns('*.docx','*.pdf','~$*'))
    (destination/'README.md').write_text('# CafeLink source\n\nBuild/run/test: [cafelink-system/README.md](cafelink-system/README.md).\n\nFinal report: [report/CafeLink_Final_Report.md](report/CafeLink_Final_Report.md). The sealed E0 ZIP in dist is needed only for the baseline integrity guard.\n',encoding='utf-8')
    forbidden=[p for p in destination.rglob('*') if p.is_file() and (p.suffix in ['.db','.sqlite','.sqlite3','.key'] or '.qr-key' in p.name)]
    if forbidden:raise ValueError('Runtime data accidentally included')
    print(json.dumps({'source':str(destination),'files':sum(p.is_file() for p in destination.rglob('*'))}))

def refresh_reports(source):
    source=Path(source);report=source/'report';report.mkdir(exist_ok=True)
    for name in ['CafeLink_Final_Report.md','SUBMISSION_READINESS.md']:
        shutil.copy2(ROOT/'report'/name,report/name)
    for name in ['tools','templates','assets/cafelink-final']:
        shutil.copytree(ROOT/'report'/name,report/name,dirs_exist_ok=True,ignore=IGNORE)
    for rel in ['cafelink-final','cafelink-v2/v9.3-gate/20261025-184703-1bdff3d4','cafelink-v2/v9.3-package/20261025-release-01']:
        shutil.copytree(ROOT/'report/evidence'/rel,report/'evidence'/rel,dirs_exist_ok=True,ignore=shutil.ignore_patterns('*.db','*.key','*.zip','*.docx','*.pdf','__pycache__'))

def validate_delivery(destination):
    from docx import Document
    import fitz
    critical=['README.md','SUBMISSION_READINESS.md','Source/README.md','Source/cafelink-system/README.md','Source/cafelink-system/docs/FINAL_ARCHITECTURE.md','Source/cafelink-system/docs/FINAL_PROTOCOL.md','Source/cafelink-system/docs/FINAL_RUNBOOK.md','Source/report/CafeLink_Final_Report.md','Source/report/SUBMISSION_READINESS.md','Report/CafeLink_Final_Report.md','Report/SUBMISSION_READINESS.md']
    links=0
    for name in critical:
        p=destination/name;text=p.read_text(encoding='utf-8-sig')
        if '\ufffd' in text:raise ValueError('Broken Unicode in '+name)
        for ref in re.findall(r'\[[^\]]*\]\(([^)]+)\)',text):
            ref=ref.split('#',1)[0].strip('<>')
            if not ref or re.match(r'^[a-zA-Z]+:',ref):continue
            path=(p.parent/ref).resolve()
            if not path.is_relative_to(destination) or not path.exists():raise ValueError('Broken or external local link: '+name+' -> '+ref)
            links+=1
    bad=[p for p in (destination/'Source').rglob('*') if p.is_file() and (p.suffix in ['.db','.sqlite','.key'] or '.qr-key' in p.name)]
    if bad:raise ValueError('Live data found in exported source')
    for name in ['node_modules','target','.tools']:
        if (destination/'Source/cafelink-system'/name).exists():raise ValueError('Developer artifacts included')
    doc=Document(destination/'Report/NhomXX_CafeLink_Report.docx')
    pdf=fitz.open(destination/'Report/NhomXX_CafeLink_Report.pdf')
    body=unicodedata.normalize('NFC','\n'.join(p.get_text() for p in pdf))
    for title in ['Abstract','Kết luận','Tài liệu tham khảo','1.348']:
        if title not in body:raise ValueError('Report content missing: '+title)
    if len(doc.tables)!=9 or len(doc.inline_shapes)!=5:raise ValueError('DOCX table/image count changed')
    runtime=destination/'Run/CafeLink-Web-V9.3'
    rm=json.loads((runtime/'SHA256-MANIFEST.json').read_text(encoding='utf-8-sig'))
    for row in rm['files']:
        if digest(runtime/row['path'])!=row['sha256']:raise ValueError('Runtime manifest mismatch')
    return {'status':'PASS','criticalDocuments':len(critical),'localLinksChecked':links,'sourceHasLiveData':False,'sourceHasDeveloperCaches':False,'wordPdfPages':len(pdf),'docxTables':len(doc.tables),'docxImages':len(doc.inline_shapes),'runtimeManifestFiles':len(rm['files']),'physicalPhone':'NOT_RUN','compilatio':'PENDING','groupMetadata':'PENDING','officialSubmission':'NOT_SUBMITTED'}

def bundle(destination):
    destination=Path(destination).resolve()
    if destination.exists():raise FileExistsError('Select a fresh final bundle')
    destination.mkdir(parents=True);copy_source(destination/'Source');refresh_reports(destination/'Source')
    report=destination/'Report';report.mkdir()
    for suffix in ['docx','pdf']:
        shutil.copy2(ROOT/f'report/NhomXX_CafeLink_Report.{suffix}',report/f'NhomXX_CafeLink_Report.{suffix}')
    text=(ROOT/'report/CafeLink_Final_Report.md').read_text(encoding='utf-8')
    text=text.replace('](assets/','](../Source/report/assets/').replace('](evidence/','](../Source/report/evidence/').replace('](../cafelink-system/','](../Source/cafelink-system/')
    (report/'CafeLink_Final_Report.md').write_text(text,encoding='utf-8')
    readiness=(ROOT/'report/SUBMISSION_READINESS.md').read_text(encoding='utf-8')
    root_ready=readiness.replace('](evidence/','](Source/report/evidence/').replace('](CafeLink_Final_Report.md)','](Report/CafeLink_Final_Report.md)').replace('](../final-project/','](Source/final-project/')
    (destination/'SUBMISSION_READINESS.md').write_text(root_ready,encoding='utf-8')
    report_ready=readiness.replace('](evidence/','](../Source/report/evidence/').replace('](../final-project/','](../Source/final-project/')
    (report/'SUBMISSION_READINESS.md').write_text(report_ready,encoding='utf-8')
    run=destination/'Run/CafeLink-Web-V9.3';run.mkdir(parents=True)
    archive=ROOT/'dist/CafeLink-Web-V9.3.zip'
    if digest(archive)!='FDCA5E041BCB83F61260056BB8B6DB859BB513EE19A8F626F4BD464F77C83CA6':raise ValueError('Frozen V9.3 ZIP changed')
    with zipfile.ZipFile(archive) as z:z.extractall(run)
    shutil.copy2(ROOT/'report/templates/Submission_README.md',destination/'README.md')
    validation=validate_delivery(destination)
    (destination/'DELIVERY_VALIDATION.json').write_text(json.dumps(validation,indent=2)+'\n',encoding='utf-8')
    manifest=[]
    for p in sorted(destination.rglob('*')):
        if p.is_file():manifest.append({'path':p.relative_to(destination).as_posix(),'bytes':p.stat().st_size,'sha256':digest(p)})
    (destination/'SHA256-MANIFEST.json').write_text(json.dumps({'format':'CafeLinkSubmission/1','createdAtUtc':datetime.now(timezone.utc).isoformat(),'files':manifest},indent=2)+'\n',encoding='utf-8')
    for out,base,prefix in [(destination.with_suffix('.zip'),destination,destination.name),(destination.parent/'NhomXX_CafeLink_Source.zip',destination/'Source','NhomXX_CafeLink_Source')]:
        if out.exists():raise FileExistsError('Do not replace an existing archive')
        with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
            for p in sorted(base.rglob('*')):
                if p.is_file():z.write(p,Path(prefix)/p.relative_to(base))
        print(json.dumps({'archive':str(out),'sha256':digest(out),'bytes':out.stat().st_size}))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--source');parser.add_argument('--bundle');args=parser.parse_args()
    if args.source:copy_source(args.source)
    elif args.bundle:bundle(args.bundle)
    else:parser.error('Choose --source or --bundle')
